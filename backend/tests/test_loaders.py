from docx import Document as DocxDocument
from pypdf import PdfWriter

import pytest

from app.core.exceptions import UnsupportedFileError
from app.rag.loaders import DocumentLoader, sanitize_filename, validate_mime_type


def test_pdf_and_docx_loaders_basic(tmp_path) -> None:
    pdf_path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with pdf_path.open("wb") as stream:
        writer.write(stream)
    parsed_pdf = DocumentLoader().load(pdf_path)
    assert parsed_pdf.source_file == "blank.pdf"
    assert parsed_pdf.sections == []

    docx_path = tmp_path / "derivative.docx"
    document = DocxDocument()
    document.add_heading("导数", level=1)
    document.add_paragraph("导数表示瞬时变化率。")
    document.save(docx_path)
    parsed_docx = DocumentLoader().load(docx_path)
    assert len(parsed_docx.sections) == 1
    assert parsed_docx.sections[0].chapter == "导数"
    assert "瞬时变化率" in parsed_docx.sections[0].text


def test_filename_sanitization_and_mime_validation() -> None:
    assert sanitize_filename("../../秘密 资料.md") == "秘密_资料.md"
    with pytest.raises(UnsupportedFileError):
        validate_mime_type("notes.pdf", "text/plain")
