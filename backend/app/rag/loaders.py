from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.exceptions import DocumentProcessingError, UnsupportedFileError


ALLOWED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}
ALLOWED_MIME_TYPES: dict[str, set[str]] = {
    ".md": {"text/markdown", "text/x-markdown", "text/plain", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
        "application/octet-stream",
    },
}


@dataclass(frozen=True, slots=True)
class ParsedSection:
    text: str
    title: str = ""
    chapter: str = ""
    section: str = ""
    page: int | None = None
    order: int = 0


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    source_file: str
    sections: list[ParsedSection]


def display_filename(filename: str) -> str:
    normalized = unicodedata.normalize("NFKC", filename or "").replace("\\", "/")
    basename = normalized.rsplit("/", 1)[-1].strip()
    basename = "".join(character for character in basename if character.isprintable())
    if not basename:
        raise UnsupportedFileError("文件名无效")
    return basename[:255]


def sanitize_filename(filename: str) -> str:
    basename = display_filename(filename)
    basename = re.sub(r"[^\w\-.()\u4e00-\u9fff]+", "_", basename, flags=re.UNICODE)
    basename = basename.lstrip(".")
    if not basename:
        raise UnsupportedFileError("文件名无效")
    suffix = Path(basename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileError("仅支持 .md、.txt、.pdf 和 .docx 文件")
    if len(basename) > 180:
        basename = f"{Path(basename).stem[: 180 - len(suffix)]}{suffix}"
    return basename


def validate_mime_type(filename: str, content_type: str | None) -> None:
    suffix = Path(filename).suffix.lower()
    mime = (content_type or "application/octet-stream").split(";", 1)[0].strip().lower()
    if suffix not in ALLOWED_MIME_TYPES or mime not in ALLOWED_MIME_TYPES[suffix]:
        raise UnsupportedFileError(
            f"文件扩展名 {suffix or '(无)'} 与 MIME 类型 {mime or '(无)'} 不匹配"
        )


class DocumentLoader:
    def load(self, path: str | Path, *, source_file: str | None = None) -> ParsedDocument:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix not in ALLOWED_EXTENSIONS:
            raise UnsupportedFileError()
        try:
            if suffix == ".md":
                sections = self._load_markdown(file_path)
            elif suffix == ".txt":
                sections = self._load_text(file_path)
            elif suffix == ".pdf":
                sections = self._load_pdf(file_path)
            else:
                sections = self._load_docx(file_path)
        except (DocumentProcessingError, UnsupportedFileError):
            raise
        except Exception as exc:
            raise DocumentProcessingError(
                internal_message=f"解析 {file_path.name} 失败: {type(exc).__name__}: {exc}"
            ) from exc
        return ParsedDocument(source_file=source_file or file_path.name, sections=sections)

    @staticmethod
    def _read_text(path: Path) -> str:
        data = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise DocumentProcessingError("文本编码无法识别，请使用 UTF-8 编码")

    def _load_markdown(self, path: Path) -> list[ParsedSection]:
        text = self._read_text(path)
        lines = text.splitlines()
        sections: list[ParsedSection] = []
        buffer: list[str] = []
        chapter = ""
        section = ""
        title = ""
        in_fence = False

        def flush() -> None:
            content = "\n".join(buffer).strip()
            if content:
                sections.append(
                    ParsedSection(
                        text=content,
                        title=title,
                        chapter=chapter,
                        section=section,
                        order=len(sections),
                    )
                )
            buffer.clear()

        for line in lines:
            stripped = line.strip()
            if stripped.startswith("```") or stripped.startswith("~~~"):
                in_fence = not in_fence
            heading = None if in_fence else re.match(r"^(#{1,3})\s+(.+?)\s*$", line)
            if heading:
                flush()
                level = len(heading.group(1))
                heading_text = heading.group(2).strip().rstrip("#").strip()
                title = heading_text
                if level == 1:
                    chapter = heading_text
                    section = ""
                else:
                    section = heading_text
                buffer.append(line)
            else:
                buffer.append(line)
        flush()
        return sections

    def _load_text(self, path: Path) -> list[ParsedSection]:
        text = self._read_text(path)
        paragraphs = re.split(r"\n\s*\n+", text)
        return [
            ParsedSection(text=paragraph.strip(), order=index)
            for index, paragraph in enumerate(paragraphs)
            if paragraph.strip()
        ]

    @staticmethod
    def _load_pdf(path: Path) -> list[ParsedSection]:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentProcessingError("PDF 已加密，无法解析") from exc
        sections: list[ParsedSection] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(
                    ParsedSection(text=text, page=page_number, order=len(sections))
                )
        return sections

    @staticmethod
    def _load_docx(path: Path) -> list[ParsedSection]:
        document = DocxDocument(str(path))
        sections: list[ParsedSection] = []
        buffer: list[str] = []
        chapter = ""
        section = ""
        title = ""

        def flush() -> None:
            content = "\n\n".join(buffer).strip()
            if content:
                sections.append(
                    ParsedSection(
                        text=content,
                        title=title,
                        chapter=chapter,
                        section=section,
                        order=len(sections),
                    )
                )
            buffer.clear()

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower()
            heading = re.match(r"heading\s*([1-3])", style_name)
            if heading:
                flush()
                level = int(heading.group(1))
                title = text
                if level == 1:
                    chapter = text
                    section = ""
                else:
                    section = text
                buffer.append(text)
            else:
                buffer.append(text)
        flush()
        return sections
